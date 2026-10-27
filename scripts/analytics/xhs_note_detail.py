"""Parse 小红书单篇「数据明细表」xlsx and merge into notes CSV rows."""

from __future__ import annotations

import argparse
import json
import re
import sys
from pathlib import Path

ROOT = Path(__file__).resolve().parents[2]
sys.path.insert(0, str(ROOT / "scripts" / "analytics"))

from xhs_note_links import normalize_title  # noqa: E402

OVERVIEW_SHEET = "基础数据总览"


def _norm_label(raw: object) -> str:
    return re.sub(r"\s+", "", str(raw or "").strip())


# Exact labels only — avoid matching「曝光数粉丝占比」as 曝光数
LABEL_TO_FIELD = {
    _norm_label("曝光数"): "impressions",
    _norm_label("观看数"): "views",
    _norm_label("封面点击率(%)"): "cover_ctr_pct",
    _norm_label("平均观看时长(s)"): "avg_watch_sec",
    _norm_label("完播率(%)"): "completion_rate_pct",
    _norm_label("2秒退出率(%)"): "exit_2s_rate_pct",
    _norm_label("涨粉数"): "follows_from_note",
}


def title_from_detail_filename(path: Path) -> str:
    name = path.name
    if name.endswith("-数据明细表.xlsx"):
        return name[: -len("-数据明细表.xlsx")]
    return path.stem


def parse_overview(path: Path) -> dict[str, str]:
    try:
        import openpyxl
    except ImportError:
        raise SystemExit("error: pip install openpyxl") from None

    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    if OVERVIEW_SHEET not in wb.sheetnames:
        wb.close()
        raise ValueError(f"{path}: missing sheet {OVERVIEW_SHEET!r}")
    ws = wb[OVERVIEW_SHEET]
    rows = list(ws.iter_rows(values_only=True))
    wb.close()
    out: dict[str, str] = {"detail_source": str(path.name)}
    for row in rows[1:]:
        if not row or row[0] is None:
            continue
        field = LABEL_TO_FIELD.get(_norm_label(row[0]))
        if not field:
            continue
        val = row[1]
        if val is None:
            continue
        if isinstance(val, float) and field.endswith("_pct"):
            out[field] = str(val)
        elif isinstance(val, float):
            out[field] = str(int(val) if val == int(val) else val)
        else:
            out[field] = str(val).strip()
    return out


def parse_daily_series(path: Path, sheet_suffix: str) -> list[dict[str, str]]:
    try:
        import openpyxl
    except ImportError:
        return []

    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    out: list[dict[str, str]] = []
    for name in wb.sheetnames:
        if not name.endswith(sheet_suffix):
            continue
        ws = wb[name]
        rows = list(ws.iter_rows(values_only=True))
        if len(rows) < 2:
            continue
        for row in rows[1:]:
            if not row or row[0] is None:
                continue
            out.append({"date": str(row[0]).strip(), "value": str(row[1]).strip()})
    wb.close()
    return out


def parse_daily_series_for(path: Path, sheet_name: str) -> list[dict[str, str]]:
    try:
        import openpyxl
    except ImportError:
        return []

    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    if sheet_name not in wb.sheetnames:
        wb.close()
        return []
    ws = wb[sheet_name]
    rows = list(ws.iter_rows(values_only=True))
    wb.close()
    series: list[dict[str, str]] = []
    for row in rows[1:]:
        if not row or row[0] is None:
            continue
        series.append({"date": str(row[0]).strip(), "value": str(row[1]).strip()})
    return series


def full_detail_json(path: Path) -> dict:
    title = title_from_detail_filename(path)
    overview = parse_overview(path)
    return {
        "title": title,
        "title_norm": normalize_title(title),
        "overview": overview,
        "series": {
            "impressions_daily": parse_daily_series_for(path, "曝光数（天）"),
            "views_daily": parse_daily_series_for(path, "观看数（天）"),
            "cover_ctr_daily": parse_daily_series_for(path, "封面点击率（天）"),
            "avg_watch_daily": parse_daily_series_for(path, "平均观看时长（天）"),
            "completion_daily": parse_daily_series_for(path, "完播率（天）"),
            "exit_2s_daily": parse_daily_series_for(path, "2s退出率（天）"),
            "follows_daily": parse_daily_series_for(path, "涨粉数（天）"),
        },
    }


def load_details_from_dir(directory: Path) -> dict[str, dict[str, str]]:
    """Map normalized title → overview fields to merge into CSV."""
    by_title: dict[str, dict[str, str]] = {}
    if not directory.is_dir():
        return by_title
    for path in sorted(directory.glob("*数据明细表.xlsx")):
        try:
            bundle = full_detail_json(path)
        except ValueError as exc:
            print(f"skip {path.name}: {exc}", file=sys.stderr)
            continue
        overview = bundle.get("overview") or {}
        if not overview:
            continue
        by_title[bundle["title_norm"]] = overview
    return by_title


def merge_details_into_rows(
    rows: list[dict[str, str]], by_title: dict[str, dict[str, str]]
) -> tuple[list[dict[str, str]], int]:
    merged = 0
    out: list[dict[str, str]] = []
    for row in rows:
        r = dict(row)
        nt = normalize_title(r.get("title") or "")
        hit = by_title.get(nt)
        if not hit:
            # fuzzy: detail filename may omit emoji suffix variants
            for key, detail in by_title.items():
                if key in nt or nt in key:
                    hit = detail
                    break
        if hit:
            for k, v in hit.items():
                if k == "detail_source":
                    r["detail_source"] = v
                elif k in ("impressions", "views", "follows_from_note"):
                    if v:
                        r[k] = v
                elif v:
                    r[k] = v
            merged += 1
        out.append(r)
    return out, merged


def cmd_ingest(directory: Path, archive_dir: Path, json_dir: Path) -> int:
    import shutil

    archive_dir.mkdir(parents=True, exist_ok=True)
    json_dir.mkdir(parents=True, exist_ok=True)
    count = 0
    for path in sorted(directory.glob("*数据明细表.xlsx")):
        dest = archive_dir / path.name
        shutil.copy2(path, dest)
        data = full_detail_json(path)
        slug = re.sub(r"[^\w\u4e00-\u9fff]+", "-", data["title"][:40]).strip("-") or "note"
        json_path = json_dir / f"{slug}.json"
        json_path.write_text(json.dumps(data, ensure_ascii=False, indent=2), encoding="utf-8")
        print(f"archived {path.name} → {json_path.relative_to(ROOT)}")
        count += 1
    print(f"ingested {count} detail workbook(s)")
    return 0 if count else 1


def main() -> int:
    parser = argparse.ArgumentParser()
    sub = parser.add_subparsers(dest="cmd", required=True)
    p = sub.add_parser("ingest", help="Copy detail xlsx + write per-note JSON series")
    p.add_argument("--dir", type=Path, required=True, help="Folder with *数据明细表.xlsx")
    p.add_argument(
        "--archive",
        type=Path,
        default=ROOT / "content/analytics/raw/xiaohongshu/note-details",
    )
    p.add_argument(
        "--json-out",
        type=Path,
        default=ROOT / "content/analytics/xiaohongshu/detail-json",
    )
    args = parser.parse_args()
    d = args.dir if args.dir.is_absolute() else (ROOT / args.dir).resolve()
    archive = args.archive if args.archive.is_absolute() else (ROOT / args.archive).resolve()
    json_out = args.json_out if args.json_out.is_absolute() else (ROOT / args.json_out).resolve()
    if args.cmd == "ingest":
        return cmd_ingest(d, archive, json_out)
    return 2


if __name__ == "__main__":
    sys.exit(main())
