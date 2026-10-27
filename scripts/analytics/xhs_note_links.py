"""Merge Xiaohongshu note URLs into analytics rows (note_id = link)."""

from __future__ import annotations

import argparse
import csv
import re
import sys
from pathlib import Path

ROOT = Path(__file__).resolve().parents[2]
sys.path.insert(0, str(ROOT / "scripts" / "analytics"))
DEFAULT_REGISTRY = ROOT / "content/analytics/xiaohongshu/note-links.csv"

REGISTRY_FIELDS = ("publish_date", "title", "note_url", "brief_path")
URL_RE = re.compile(r"^https?://", re.I)

WORKBOOK_SHEET = "笔记链接回填"
WORKBOOK_HEADERS = (
    "发布日期",
    "笔记标题",
    "曝光",
    "观看量",
    "点赞",
    "收藏",
    "评论",
    "分享",
    "涨粉",
    "支柱",
    "笔记链接（请填写）",
    "Brief路径（可选）",
)
DEFAULT_WORKBOOK = ROOT / "content/analytics/xiaohongshu/笔记链接回填表.xlsx"
PUBLISHED_NOTES_INDEX = ROOT / "content/xiaohongshu/published-notes-index.md"
BRIEF_PATH_INDEX = "content/xiaohongshu/published-notes-index.md"


def normalize_title(title: str) -> str:
    return re.sub(r"\s+", "", (title or "").strip())


def _registry_key(publish_date: str, title: str) -> tuple[str, str]:
    return (publish_date.strip(), normalize_title(title))


def load_registry(path: Path) -> list[dict[str, str]]:
    if not path.is_file():
        return []
    with path.open(encoding="utf-8-sig", newline="") as f:
        rows = list(csv.DictReader(f))
    out: list[dict[str, str]] = []
    for row in rows:
        out.append(
            {
                "publish_date": (row.get("publish_date") or "").strip(),
                "title": (row.get("title") or "").strip(),
                "note_url": (row.get("note_url") or "").strip(),
                "brief_path": (row.get("brief_path") or "").strip(),
            }
        )
    return out


def registry_lookup(registry: list[dict[str, str]]) -> tuple[dict[tuple[str, str], dict[str, str]], dict[str, dict[str, str]]]:
    by_date_title: dict[tuple[str, str], dict[str, str]] = {}
    by_title: dict[str, dict[str, str]] = {}
    for rec in registry:
        if not rec.get("title"):
            continue
        key = _registry_key(rec["publish_date"], rec["title"])
        by_date_title[key] = rec
        nt = normalize_title(rec["title"])
        if nt not in by_title:
            by_title[nt] = rec
    return by_date_title, by_title


def resolve_link(
    publish_date: str,
    title: str,
    by_date_title: dict[tuple[str, str], dict[str, str]],
    by_title: dict[str, dict[str, str]],
) -> dict[str, str] | None:
    key = _registry_key(publish_date, title)
    rec = by_date_title.get(key)
    if rec and rec.get("note_url") and URL_RE.match(rec["note_url"]):
        return rec
    nt = normalize_title(title)
    rec = by_title.get(nt)
    if rec and rec.get("note_url") and URL_RE.match(rec["note_url"]):
        return rec
    return None


def apply_links_to_rows(
    rows: list[dict[str, str]],
    registry: list[dict[str, str]],
) -> tuple[list[dict[str, str]], int]:
    by_date_title, by_title = registry_lookup(registry)
    applied = 0
    out: list[dict[str, str]] = []
    for row in rows:
        r = dict(row)
        pub = r.get("publish_date") or ""
        title = r.get("title") or ""
        hit = resolve_link(pub, title, by_date_title, by_title)
        if hit:
            r["note_id"] = hit["note_url"]
            if hit.get("brief_path"):
                r["brief_path"] = hit["brief_path"]
            applied += 1
        out.append(r)
    return out, applied


def write_registry(path: Path, rows: list[dict[str, str]]) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    with path.open("w", encoding="utf-8", newline="") as f:
        w = csv.DictWriter(f, fieldnames=list(REGISTRY_FIELDS))
        w.writeheader()
        for row in rows:
            w.writerow({k: row.get(k, "") for k in REGISTRY_FIELDS})


def merge_registry(existing: list[dict[str, str]], incoming: list[dict[str, str]]) -> list[dict[str, str]]:
    """Keep existing note_url/brief_path when incoming leaves them blank."""
    index: dict[tuple[str, str], dict[str, str]] = {}
    for rec in existing:
        if rec.get("title"):
            index[_registry_key(rec["publish_date"], rec["title"])] = dict(rec)
    for rec in incoming:
        if not rec.get("title"):
            continue
        key = _registry_key(rec["publish_date"], rec["title"])
        prev = index.get(key, {})
        merged = {
            "publish_date": rec["publish_date"] or prev.get("publish_date", ""),
            "title": rec["title"],
            "note_url": rec.get("note_url") or prev.get("note_url", ""),
            "brief_path": rec.get("brief_path") or prev.get("brief_path", ""),
        }
        index[key] = merged
    return sorted(index.values(), key=lambda r: (r["publish_date"], r["title"]))


def infer_brief_path(title: str, xhs_dir: Path) -> str:
    """Topic Manager Brief if title matches; else published-notes index."""
    nt = normalize_title(title)
    for path in sorted(xhs_dir.glob("*-brief.md")):
        text = path.read_text(encoding="utf-8")
        candidates = [_brief_title(text)]
        if m := re.search(r'"subject"\s*:\s*"([^"]+)"', text):
            candidates.append(m.group(1).strip())
        for cand in candidates:
            if cand and normalize_title(cand) == nt:
                return f"content/xiaohongshu/{path.name}"
            if cand and (nt in normalize_title(cand) or normalize_title(cand) in nt):
                return f"content/xiaohongshu/{path.name}"
    return BRIEF_PATH_INDEX


def refresh_published_notes_index(registry: list[dict[str, str]]) -> None:
    lines = [
        "# 已发布笔记索引（非 Topic Manager Brief）",
        "",
        "认知期深度 Brief 见 `2026-10-17-discovery-*`；下表为 **历史已发笔记**（艺术专业 / 日常），供 `brief_path` 与数据分析对齐。",
        "",
        "| 发布日期 | 标题 | 笔记链接 | brief_path |",
        "|----------|------|----------|------------|",
    ]
    for rec in sorted(registry, key=lambda r: (r.get("publish_date", ""), r.get("title", ""))):
        url = rec.get("note_url") or "—"
        bp = rec.get("brief_path") or BRIEF_PATH_INDEX
        title = (rec.get("title") or "").replace("|", "\\|")
        lines.append(f"| {rec.get('publish_date', '')} | {title} | {url} | `{bp}` |")
    lines.extend(["", "---", "由 `xhs_note_links.py import-workbook` 刷新。"])
    PUBLISHED_NOTES_INDEX.parent.mkdir(parents=True, exist_ok=True)
    PUBLISHED_NOTES_INDEX.write_text("\n".join(lines) + "\n", encoding="utf-8")


def _brief_title(text: str) -> str:
    if m := re.search(r"## 标题[^\n]*\n+\s*(.+?)\s*\n", text):
        return m.group(1).strip()
    if m := re.search(r'"subject"\s*:\s*"([^"]+)"', text):
        return m.group(1).strip()
    return ""


def links_from_briefs(xhs_dir: Path) -> list[dict[str, str]]:
    rows: list[dict[str, str]] = []
    for path in sorted(xhs_dir.glob("*-brief.md")):
        text = path.read_text(encoding="utf-8")
        url = ""
        pub = ""
        if m := re.search(r"发布日期[：:]\s*(\d{4}-\d{2}-\d{2})", text):
            pub = m.group(1)
        if u := re.search(r"笔记链接[：:]\s*(https?://\S+)", text):
            url = u.group(1).rstrip(")")
        if not url:
            continue
        title = _brief_title(text)
        rows.append(
            {
                "publish_date": pub,
                "title": title,
                "note_url": url,
                "brief_path": f"content/xiaohongshu/{path.name}",
            }
        )
    return rows


def _rows_for_workbook(xlsx: Path, registry: list[dict[str, str]]) -> list[dict[str, str]]:
    from xhs_convert_notes_detail import read_notes_sheet

    by_date_title, by_title = registry_lookup(registry)
    rows = read_notes_sheet(xlsx)
    out: list[dict[str, str]] = []
    for r in rows:
        hit = resolve_link(r["publish_date"], r["title"], by_date_title, by_title)
        out.append(
            {
                "publish_date": r["publish_date"],
                "title": r["title"],
                "impressions": r.get("impressions", ""),
                "views": r.get("views", ""),
                "likes": r.get("likes", ""),
                "collects": r.get("collects", ""),
                "comments": r.get("comments", ""),
                "shares": r.get("shares", ""),
                "follows_from_note": r.get("follows_from_note", ""),
                "pillar": r.get("pillar", ""),
                "note_url": (hit or {}).get("note_url", ""),
                "brief_path": (hit or {}).get("brief_path", "") or r.get("brief_path", ""),
            }
        )
    return out


def write_workbook(path: Path, data_rows: list[dict[str, str]]) -> None:
    try:
        import openpyxl
        from openpyxl.styles import Font
        from openpyxl.utils import get_column_letter
    except ImportError:
        raise SystemExit("error: pip install openpyxl") from None

    path.parent.mkdir(parents=True, exist_ok=True)
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = WORKBOOK_SHEET
    ws.append(list(WORKBOOK_HEADERS))
    for cell in ws[1]:
        cell.font = Font(bold=True)
    for r in data_rows:
        ws.append(
            [
                r["publish_date"],
                r["title"],
                r.get("impressions", ""),
                r.get("views", ""),
                r.get("likes", ""),
                r.get("collects", ""),
                r.get("comments", ""),
                r.get("shares", ""),
                r.get("follows_from_note", ""),
                r.get("pillar", ""),
                r.get("note_url", ""),
                r.get("brief_path", ""),
            ]
        )
    widths = (12, 42, 10, 10, 8, 8, 8, 8, 8, 12, 48, 36)
    for i, w in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(i)].width = w
    ws.freeze_panes = "A2"
    wb.save(path)


def read_workbook(path: Path) -> list[dict[str, str]]:
    try:
        import openpyxl
    except ImportError:
        raise SystemExit("error: pip install openpyxl") from None

    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    ws = wb[WORKBOOK_SHEET] if WORKBOOK_SHEET in wb.sheetnames else wb.active
    rows = list(ws.iter_rows(values_only=True))
    wb.close()
    if not rows:
        return []
    header = [str(c).strip() if c is not None else "" for c in rows[0]]
    col = {name: idx for idx, name in enumerate(header)}

    def cell(row: tuple, name: str) -> str:
        idx = col.get(name)
        if idx is None or idx >= len(row) or row[idx] is None:
            return ""
        return str(row[idx]).strip()

    out: list[dict[str, str]] = []
    for row in rows[1:]:
        if not row or not row[0]:
            continue
        title = cell(row, "笔记标题")
        if not title:
            continue
        out.append(
            {
                "publish_date": cell(row, "发布日期"),
                "title": title,
                "note_url": cell(row, "笔记链接（请填写）"),
                "brief_path": cell(row, "Brief路径（可选）"),
            }
        )
    return out


def scaffold_from_xlsx(xlsx: Path) -> list[dict[str, str]]:
    # Local import avoids circular dependency at module load.
    from xhs_convert_notes_detail import read_notes_sheet

    rows = read_notes_sheet(xlsx)
    return [
        {
            "publish_date": r["publish_date"],
            "title": r["title"],
            "note_url": "",
            "brief_path": r.get("brief_path") or "",
        }
        for r in rows
    ]


def cmd_scaffold(registry: Path, xlsx: Path) -> int:
    if not xlsx.is_file():
        print(f"error: not found {xlsx}", file=sys.stderr)
        return 2
    incoming = scaffold_from_xlsx(xlsx)
    merged = merge_registry(load_registry(registry), incoming)
    write_registry(registry, merged)
    filled = sum(1 for r in merged if r.get("note_url"))
    print(f"registry → {registry.relative_to(ROOT)} ({len(merged)} rows, {filled} with URL)")
    return 0


def cmd_sync_briefs(registry: Path, xhs_dir: Path) -> int:
    incoming = links_from_briefs(xhs_dir)
    if not incoming:
        print("no briefs with 笔记链接 found")
    merged = merge_registry(load_registry(registry), incoming)
    write_registry(registry, merged)
    filled = sum(1 for r in merged if r.get("note_url"))
    print(f"merged {len(incoming)} from briefs → {len(merged)} rows ({filled} with URL)")
    return 0


def _read_analytics_csv(path: Path) -> tuple[list[str], list[dict[str, str]]]:
    with path.open(encoding="utf-8-sig", newline="") as f:
        reader = csv.DictReader(f)
        fieldnames = list(reader.fieldnames or [])
        rows = list(reader)
    return fieldnames, rows


def cmd_export_workbook(registry: Path, xlsx: Path, out: Path) -> int:
    if not xlsx.is_file():
        print(f"error: not found {xlsx}", file=sys.stderr)
        return 2
    data = _rows_for_workbook(xlsx, load_registry(registry))
    write_workbook(out, data)
    filled = sum(1 for r in data if r.get("note_url") and URL_RE.match(r["note_url"]))
    print(f"workbook → {out.relative_to(ROOT)} ({len(data)} rows, {filled} links already filled)")
    return 0


def cmd_import_workbook(registry: Path, workbook: Path, xhs_dir: Path) -> int:
    if not workbook.is_file():
        print(f"error: not found {workbook}", file=sys.stderr)
        return 2
    incoming = read_workbook(workbook)
    merged = merge_registry(load_registry(registry), incoming)
    for rec in merged:
        if not rec.get("brief_path"):
            rec["brief_path"] = infer_brief_path(rec.get("title") or "", xhs_dir)
    write_registry(registry, merged)
    refresh_published_notes_index(merged)
    filled = sum(1 for r in merged if r.get("note_url") and URL_RE.match(r["note_url"]))
    print(
        f"imported workbook → {registry.relative_to(ROOT)} "
        f"({len(incoming)} rows read, {filled} with URL)"
    )
    print(f"index → {PUBLISHED_NOTES_INDEX.relative_to(ROOT)}")
    return 0


def cmd_apply(registry: Path, csv_path: Path) -> int:
    if not csv_path.is_file():
        print(f"error: not found {csv_path}", file=sys.stderr)
        return 2
    registry_rows = load_registry(registry)
    if not registry_rows:
        print(f"warn: empty registry {registry}", file=sys.stderr)
    fieldnames, rows = _read_analytics_csv(csv_path)
    if "note_id" not in fieldnames:
        print("error: CSV missing note_id column", file=sys.stderr)
        return 1
    updated, applied = apply_links_to_rows(rows, registry_rows)
    with csv_path.open("w", encoding="utf-8", newline="") as f:
        w = csv.DictWriter(f, fieldnames=fieldnames)
        w.writeheader()
        w.writerows(updated)
    print(f"applied {applied} links → {csv_path.relative_to(ROOT)}")
    return 0


def main(argv: list[str] | None = None) -> int:
    parser = argparse.ArgumentParser(description="Xiaohongshu note URL registry")
    parser.add_argument(
        "--registry",
        type=Path,
        default=DEFAULT_REGISTRY,
        help=f"Link registry CSV (default: {DEFAULT_REGISTRY.relative_to(ROOT)})",
    )
    sub = parser.add_subparsers(dest="cmd", required=True)

    p_scaffold = sub.add_parser("scaffold", help="Build/merge registry rows from notes-detail xlsx")
    p_scaffold.add_argument("--xlsx", type=Path, required=True)

    p_sync = sub.add_parser("sync-briefs", help="Merge URLs from Brief ## 发布状态")
    p_sync.add_argument(
        "--xhs-dir",
        type=Path,
        default=ROOT / "content/xiaohongshu",
    )

    p_apply = sub.add_parser("apply", help="Patch note_id in an analytics CSV")
    p_apply.add_argument("--csv", type=Path, required=True)

    p_export = sub.add_parser(
        "export-workbook",
        help="Excel 回填表：含后台全部指标 + 笔记链接（请填写）",
    )
    p_export.add_argument(
        "--xlsx",
        type=Path,
        required=True,
        help="notes-detail.xlsx from 创作者后台",
    )
    p_export.add_argument(
        "--out",
        type=Path,
        default=DEFAULT_WORKBOOK,
        help=f"Output .xlsx (default: {DEFAULT_WORKBOOK.relative_to(ROOT)})",
    )

    p_import = sub.add_parser(
        "import-workbook",
        help="读填好的 Excel → 更新 note-links.csv",
    )
    p_import.add_argument(
        "--workbook",
        type=Path,
        default=DEFAULT_WORKBOOK,
        help="Filled 笔记链接回填表.xlsx",
    )
    p_import.add_argument(
        "--xhs-dir",
        type=Path,
        default=ROOT / "content/xiaohongshu",
    )

    args = parser.parse_args(argv)
    registry = args.registry if args.registry.is_absolute() else (ROOT / args.registry).resolve()

    if args.cmd == "scaffold":
        xlsx = args.xlsx if args.xlsx.is_absolute() else (ROOT / args.xlsx).resolve()
        return cmd_scaffold(registry, xlsx)
    if args.cmd == "sync-briefs":
        xhs_dir = args.xhs_dir if args.xhs_dir.is_absolute() else (ROOT / args.xhs_dir).resolve()
        return cmd_sync_briefs(registry, xhs_dir)
    if args.cmd == "apply":
        csv_path = args.csv if args.csv.is_absolute() else (ROOT / args.csv).resolve()
        return cmd_apply(registry, csv_path)
    if args.cmd == "export-workbook":
        xlsx = args.xlsx if args.xlsx.is_absolute() else (ROOT / args.xlsx).resolve()
        out = args.out if args.out.is_absolute() else (ROOT / args.out).resolve()
        return cmd_export_workbook(registry, xlsx, out)
    if args.cmd == "import-workbook":
        wb = args.workbook if args.workbook.is_absolute() else (ROOT / args.workbook).resolve()
        xhs_dir = args.xhs_dir if args.xhs_dir.is_absolute() else (ROOT / args.xhs_dir).resolve()
        return cmd_import_workbook(registry, wb, xhs_dir)
    return 2


if __name__ == "__main__":
    sys.exit(main())
